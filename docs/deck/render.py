"""Render the HTML slides to PNG with headless Edge, then assemble PPTX and PDF.

    py -3.12 docs/deck/render.py            # all slides
    py -3.12 docs/deck/render.py 03 08      # only these slide numbers

Slides live in docs/deck/src/NN-name.html (1920x1080 CSS px), rendered at 2x device scale
(3840x2160) for crisp text. Output: docs/deck/out/NN.png, docs/deck/Sahaay_Pitch.pptx and
docs/deck/Sahaay_Pitch.pdf. No window is opened: Edge runs headless.
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

from PIL import Image

HERE = Path(__file__).resolve().parent
SRC = HERE / "src"
OUT = HERE / "out"
EDGE = [r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe", r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe"]


def browser() -> str:
    for p in EDGE:
        if Path(p).exists():
            return p
    raise FileNotFoundError("no headless-capable browser found")


def render(html: Path, png: Path) -> None:
    b = browser()
    cmd = [b, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--no-first-run", "--no-default-browser-check",
           "--disable-extensions", "--window-size=1920,1080", "--force-device-scale-factor=2", "--virtual-time-budget=4000",
           "--default-background-color=00000000", f"--screenshot={png}", html.resolve().as_uri()]
    subprocess.run(cmd, check=True, capture_output=True, timeout=120)
    im = Image.open(png)
    if im.size != (3840, 2160):
        im = im.resize((3840, 2160), Image.LANCZOS)
        im.save(png)


def slides() -> list[Path]:
    return sorted(p for p in SRC.glob("*.html") if re.match(r"\d\d[a-z]?-", p.name))


def notes_for(html: Path) -> str:
    """Speaker notes = the <!-- notes: ... --> comment in the slide, if any."""
    m = re.search(r"<!--\s*notes:(.*?)-->", html.read_text(encoding="utf-8"), re.S)
    return m.group(1).strip() if m else ""


def assemble(pngs: list[tuple[Path, Path]]) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for html, png in pngs:
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
        n = notes_for(html)
        if n:
            s.notes_slide.notes_text_frame.text = n
    def save_or_fallback(path: Path, saver) -> Path:
        """If the target is open in PowerPoint or a PDF viewer, write a .new copy instead of failing."""
        try:
            saver(path)
            return path
        except PermissionError:
            alt = path.with_name(path.stem + ".new" + path.suffix)
            saver(alt)
            print(f"WARNING: {path.name} is open in another program; wrote {alt.name} instead. Close it and rename.")
            return alt

    pptx_path = save_or_fallback(HERE / "Sahaay_Pitch.pptx", lambda p: prs.save(str(p)))
    # PDF: same frames, JPEG-compressed pages
    pages = []
    for _, png in pngs:
        im = Image.open(png).convert("RGB").resize((2560, 1440), Image.LANCZOS)
        pages.append(im)
    pdf_path = save_or_fallback(HERE / "Sahaay_Pitch.pdf",
                                lambda p: pages[0].save(str(p), save_all=True, append_images=pages[1:], resolution=192.0, quality=90))
    print(f"wrote {pptx_path.name} ({len(pngs)} slides, {pptx_path.stat().st_size/1e6:.1f} MB) and {pdf_path.name} ({pdf_path.stat().st_size/1e6:.1f} MB)")


def main() -> int:
    OUT.mkdir(exist_ok=True)
    only = set(sys.argv[1:])
    all_slides = slides()
    pairs = []
    for html in all_slides:
        num = html.name.split("-", 1)[0]
        png = OUT / f"{num}.png"
        if not only or num in only:
            render(html, png)
            print(f"rendered {html.name} -> {png.name}")
        pairs.append((html, png))
    if all(p.exists() for _, p in pairs):
        assemble(pairs)
    else:
        print("some slides not rendered yet; skipping assembly")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
